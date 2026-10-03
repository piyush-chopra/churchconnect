import io, re, zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from django.core.exceptions import ValidationError
SKILLS = {'Leadership':['leadership','leader','leading'], 'Teaching':['teaching','teacher','teach'], 'Youth ministry':['youth','student ministry','teen'], 'Worship':['worship','music','musician'], 'Pastoral care':['pastoral','counseling','counselling'], 'Community outreach':['outreach','community service','missions'], 'Administration':['administration','administrative','operations'], 'Communication':['communication','communications','writing'], 'Event planning':['event','events'], 'Volunteer management':['volunteer'], 'Media production':['video','media','production','audio'], 'Finance':['finance','accounting','bookkeeping'], 'Children’s ministry':['children','kids','childcare'], 'Theology':['theology','seminary','biblical'], 'Technology':['technology','software','digital']}

def parse_resume(upload):
    if upload.size > 5*1024*1024:
        raise ValidationError('Your file is larger than 5 MB. Please upload a smaller résumé.')
    suffix=Path(upload.name).suffix.lower()
    if suffix not in ['.pdf','.docx','.txt']:
        raise ValidationError('Choose a PDF, DOCX or TXT file. Legacy DOC files are not supported.')
    data=upload.read()
    try:
        if suffix=='.pdf':
            if not data.startswith(b'%PDF-'): raise ValueError()
            reader=PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages)>30: raise ValueError()
            text='\n'.join(p.extract_text() or '' for p in reader.pages)
        elif suffix=='.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist())>20*1024*1024 or 'word/document.xml' not in z.namelist(): raise ValueError()
            document=Document(io.BytesIO(data))
            text='\n'.join([p.text for p in document.paragraphs]+[c.text for t in document.tables for r in t.rows for c in r.cells])
        else:
            text=data.decode('utf-8')
            if '\x00' in text: raise ValueError()
    except Exception:
        raise ValidationError('We could not read this file. Use a text-based PDF, DOCX, or UTF-8 TXT file (up to 30 PDF pages).')
    if len(text.strip())<20: raise ValidationError('This résumé has too little readable text. For a scanned PDF, export a text-based copy first.')
    if len(text)>100000: raise ValidationError('This résumé is too long. Please upload a shorter version.')
    upload.seek(0)
    skills=[name for name,keywords in SKILLS.items() if any(re.search(r'\b'+re.escape(word)+r'\b',text.lower()) for word in keywords)]
    return text,skills
